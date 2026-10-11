#!/usr/bin/env python3
"""Verify and package the generated brand kit; no application secrets or data."""
from pathlib import Path
from datetime import datetime, timezone
import hashlib
import json
import zipfile
import xml.etree.ElementTree as E
from PIL import Image
from fontTools.ttLib import TTFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT=Path(__file__).resolve().parents[1]
B=ROOT/'public/brand'
C=json.loads((B/'brand.config.json').read_text())
P=C['palette']


def luminance(hexcode):
    vals=[int(hexcode[i:i+2],16)/255 for i in (1,3,5)]
    rgb=[v/12.92 if v<=.04045 else ((v+.055)/1.055)**2.4 for v in vals]
    return .2126*rgb[0]+.7152*rgb[1]+.0722*rgb[2]


def contrast(a,b):
    hi,lo=sorted([luminance(a),luminance(b)],reverse=True)
    return (hi+.05)/(lo+.05)


specs={}
for name,value in P.items():
    rgb=[int(value[i:i+2],16) for i in (1,3,5)]
    vals=[v/255 for v in rgb]; k=1-max(vals)
    cmy=[0,0,0] if k==1 else [(1-v-k)/(1-k) for v in vals]
    specs[name]={'hex':value,'rgb':rgb,'cmykApproximateUnprofiledPercent':[round(v*100) for v in cmy+[k]],'printNote':'Unprofiled estimate, not a press specification. Printer ICC conversion and physical proof required.'}
(B/'colour-specs.json').write_text(json.dumps(specs,indent=2)+'\n')
pairs=[('Body on warm white',P['ink'],P['cream']),('Secondary on warm white',P['muted'],P['cream']),('White on navy',P['white'],P['navy']),('White on accessible action red',P['white'],P['redAccessible']),('Blue link on warm white',P['blue'],P['cream']),('Success on white',P['green'],P['white']),('Warning on white',P['amber'],P['white']),('Dark secondary text','#CBD5E1',P['navy']),('Dark link','#93C5FD',P['navy']),('Dark success','#86EFAC',P['navy']),('Dark warning','#FCD34D',P['navy']),('Dark danger','#FF8B93',P['navy'])]
results=[]
for name,fg,bg in pairs:
    ratio=contrast(fg,bg)
    if ratio<4.5: raise ValueError(f'AA text contrast failed: {name} {ratio:.3f}')
    results.append({'pair':name,'foreground':fg,'background':bg,'ratio':round(ratio,3),'normalTextAA':True})
results.append({'pair':'Identity red / white — restricted to logo, decoration or large text','foreground':P['white'],'background':P['red'],'ratio':round(contrast(P['white'],P['red']),3),'normalTextAA':False,'largeTextAA':True})
(B/'contrast-report.json').write_text(json.dumps({'method':'WCAG 2 sRGB relative luminance','pairs':results,'scope':'Defined token pairs only; not a whole-application accessibility certification.'},indent=2)+'\n')

# Installable originals in TTF format, with the same existing OFL notices.
for filename in ['geist','inter','jetbrains-mono','noto-sans-arabic']:
    font=TTFont(B/'fonts'/(filename+'.woff2')); font.flavor=None
    font.save(B/'fonts'/(filename+'.ttf'))

# ICO is packaging of the new vector renders, not a tracing/edit of old artwork.
with Image.open(B/'favicon-48.png') as icon:
    icon.save(B/'favicon.ico',format='ICO',sizes=[(16,16),(32,32),(48,48)])
with Image.open(B/'favicon.ico') as icon:
    if icon.ico.sizes()!={(16,16),(32,32),(48,48)}: raise ValueError('ICO sizes are incomplete')
for name,size in [('apple-touch-icon.png',180),('app-icon-192.png',192),('app-icon-512.png',512),('app-icon-maskable-512.png',512),('by-mark-32.png',32),('by-mark-64.png',64)]:
    with Image.open(B/name) as im:
        if im.size!=(size,size): raise ValueError(f'Incorrect icon dimensions: {name}')

# A truly editable PowerPoint deck: native text and shapes; logo PNGs embedded.
deck=Presentation(); deck.slide_width=Inches(13.333333); deck.slide_height=Inches(7.5)
deck.core_properties.title='Loadbyton OS — Brand system 1.0'
deck.core_properties.subject='Review candidate based on Project-Loadbyton'
deck.core_properties.author='Loadbyton'
deck.core_properties.comments='Owner approval pending. No invented commercial/provider evidence.'


def rgb(hexcode): return RGBColor.from_string(hexcode.lstrip('#'))


def textbox(slide,s,x,y,w,h,size=24,color=None,font='Geist',bold=False):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame
    tf.clear();tf.word_wrap=True
    for i,line in enumerate(s.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
        p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color or P['navy'])
        p.space_after=Pt(12)
    return box


def shape(slide,x,y,w,h,color):
    s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    s.fill.solid();s.fill.fore_color.rgb=rgb(color);s.line.fill.background();return s


def slide(title,number,dark=False):
    s=deck.slides.add_slide(deck.slide_layouts[6]); bg=P['navy'] if dark else P['cream']
    s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(bg)
    logo='logo-silver.png' if dark else 'logo-navy.png'
    s.shapes.add_picture(str(B/'exports'/logo),Inches(.65),Inches(.35),width=Inches(3.0))
    textbox(s,f'{number:02d} / IDENTITY SYSTEM 1.0',.7,1.2,11.8,.35,11,P['white'] if dark else P['muted'],'JetBrains Mono')
    textbox(s,title,.7,1.8,11.6,1.5,40,P['white'] if dark else P['navy'],bold=True)
    shape(s,.7,7.02,11.8,.02,P['red'])
    textbox(s,'REVIEW CANDIDATE / OWNER APPROVAL PENDING',.7,7.08,11.8,.25,9,P['white'] if dark else P['muted'],'JetBrains Mono')
    return s

s=slide('Your logistics company,\nconnected.',1,True)
textbox(s,'Run the movement. Know the margin.',.7,4.5,10,1,24,P['white'])
s=slide('A record follows the movement.',2)
textbox(s,'Purpose\nKeep operational records, source documents and commercial context together.',.7,3.4,7.2,2,24)
textbox(s,'PRECISE\nDEPENDABLE\nDIRECT\nCONNECTED',9,3.5,3,2.8,20,P['muted'],'JetBrains Mono')
s=slide('One parent. A connected product family.',3)
textbox(s,'LOAD|BY|TON / parent identity\nLOAD|BY|TON.store / commerce and workspace\nLoadbyton OS / operating-platform shorthand',.7,3.35,11.5,2.5,25)
s=slide('Two dividers. One context.',4)
s.shapes.add_picture(str(B/'exports/logo-navy.png'),Inches(.75),Inches(3.6),width=Inches(8.5))
s.shapes.add_picture(str(B/'app-icon-192.png'),Inches(10.1),Inches(3.45),width=Inches(1.7))
textbox(s,'Use outlined SVG masters; preserve clear space and proportions.',.7,5.5,11.5,.8,20,P['muted'])
s=slide('Structure. Signal. Space.',5)
for i,(name,color) in enumerate([('Navy',P['navy']),('Red',P['red']),('Warm white',P['cream'])]):
    shape(s,.7+i*4.05,3.6,3.7,1.5,color)
    textbox(s,name+' / '+color,.7+i*4.05,5.25,3.7,.6,18)
textbox(s,'Small white labels use #C62828 for AA contrast. RGB proof files need printer ICC conversion.',.7,6.15,11.5,.7,14,P['muted'],'Inter')
s=slide('Readable operations.',6)
for i,(name,copy,font) in enumerate([('DISPLAY','Geist / 600–700','Geist'),('BODY','Inter / 400–600','Inter'),('DATA','LBT-2026-001','JetBrains Mono')]):
    textbox(s,name,.7+i*4.05,3.6,3.7,.4,11,P['muted'],'JetBrains Mono')
    textbox(s,copy,.7+i*4.05,4.3,3.7,1.3,25,P['navy'],font)
textbox(s,'Noto Sans Arabic is supplied. Arabic wording and brand transliteration still require native review.',.7,6.05,11.7,.8,17,P['muted'],'Inter')
s=slide('Say what happens.',7)
textbox(s,'Use\nReview the source document.\nAwaiting independent finance review.\nPilot availability requires confirmation.',.7,3.4,6,2.8,23)
textbox(s,'Avoid\nGuaranteed savings.\nEvery wallet works everywhere.\nOfficial clearance built in.',7.5,3.4,5,2.8,23,P['muted'])
s=slide('One system. Many surfaces.',8)
textbox(s,'Vector masters / app icons / social covers\nStationery / presentations / HTML email\nSector icons / illustrations / local fonts\nSVG + PNG + vector PDF + editable PPTX',.7,3.3,11.7,2.5,25)
s=slide('Approve the identity. Keep the gates honest.',9,True)
textbox(s,'Owner artwork acceptance / native Arabic review\nTrademark and contact verification / printer proof\nRegistrar evidence / real business and release evidence',.7,3.5,11.7,2.4,23,P['white'])
deck.save(B/'exports/loadbyton-os-brand-deck.pptx')

# Every SVG must parse; canonical masters are standalone path geometry.
svgs=list(B.rglob('*.svg'))
for p in svgs:
    root=E.parse(p).getroot()
    for e in root.iter():
        tag=e.tag.rsplit('}',1)[-1]
        if tag in ['script','image','foreignObject']: raise ValueError(f'Unsafe/raster SVG: {p.relative_to(B)} {tag}')
        if p.parent==B/'masters' or p.name.startswith('logo-') or p.name in ['by-mark.svg','favicon.svg','app-icon-maskable.svg']:
            if tag=='text': raise ValueError(f'Font-dependent master: {p.relative_to(B)}')

manifest=json.loads((B/'site.webmanifest').read_text())
for icon in manifest['icons']:
    if not (ROOT/'public'/icon['src'].lstrip('/')).exists(): raise ValueError('Manifest icon missing')
for asset in json.loads((B/'asset-index.json').read_text())['assets']:
    if not (B/asset['path']).exists(): raise ValueError('Catalogue source missing')
    if asset['editableText'] and not (B/'outlined'/asset['path']).exists(): raise ValueError('Portable template missing')

# Preserve the old PDF-extracted raster logos in the checkout, but don't call them new kit exports.
excluded={'logo-navy.png','logo-silver.png','checksums.json'}
files=[p for p in B.rglob('*') if p.is_file() and p.suffix!='.zip' and p.name not in excluded]
checksums={str(p.relative_to(B)):hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(files)}
(B/'checksums.json').write_text(json.dumps({'algorithm':'SHA-256','files':checksums},indent=2)+'\n')
archive=B/'downloads/loadbyton-os-brand-kit.zip';archive.parent.mkdir(exist_ok=True)
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED,compresslevel=9) as z:
    for p in sorted(files+[B/'checksums.json']):
        info=zipfile.ZipInfo('loadbyton-os-brand-kit/'+str(p.relative_to(B)),date_time=(2026,10,11,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED;info.external_attr=0o644<<16
        z.writestr(info,p.read_bytes())
    for script in ['build-brand.py','export-brand.mjs','package-brand.py','brand-requirements.txt']:
        z.write(ROOT/'scripts'/script,'scripts/'+script)
with zipfile.ZipFile(archive) as z:
    if z.testzip() is not None: raise ValueError('ZIP CRC check failed')
    for relative,expected in checksums.items():
        if hashlib.sha256(z.read('loadbyton-os-brand-kit/'+relative)).hexdigest()!=expected: raise ValueError('ZIP hash mismatch')
print(f'Verified {len(svgs)} SVGs, {len(pairs)} accessible text pairs, icons, catalogue, manifest and ZIP hashes.')
print(f'Packaged {len(files)+1} kit files + 4 rebuild sources: {archive.stat().st_size/1024/1024:.2f} MiB. Editable PowerPoint: 9 slides.')
